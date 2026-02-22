"""
PO packages: folder structure and Excel intrasheets for carton lists.

Each PO has a folder under data/po_packages/{po_number}/ containing:
- intrasheet.xlsx: inspection fields (carton, style, units, disposition, etc.)
- manifest.json: vendor, expected_style_count, expected_units_total (should match intrasheet)
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data" / "po_packages"


def _normalize_po(po: str) -> str:
    """Extract digits only from PO number (e.g. PO/12345 -> 12345)."""
    if not po:
        return ""
    return re.sub(r"\D", "", str(po))


def _po_folder(po: str) -> Path:
    """Path to PO package folder."""
    normalized = _normalize_po(po)
    return DATA_DIR / normalized if normalized else DATA_DIR / "_unknown"


def get_po_manifest(po: str) -> dict[str, Any]:
    """
    Load manifest.json from PO package folder if it exists.
    Returns dict with vendor, expected_units_total, expected_style_count (optional overrides).
    """
    folder = _po_folder(po)
    manifest_path = folder / "manifest.json"
    if not manifest_path.exists():
        return {}
    try:
        data = json.loads(manifest_path.read_text())
        return dict(data) if isinstance(data, dict) else {}
    except Exception:
        return {}


def _derive_from_intrasheet(cartons: list[dict[str, Any]]) -> tuple[int, int]:
    """Return (expected_style_count, expected_units_total) from carton list."""
    styles: set[str] = set()
    total = 0
    for c in cartons:
        sid = str(c.get("style_id") or c.get("Style ID") or "").strip()
        if sid:
            styles.add(sid)
        exp = c.get("expected_units") or c.get("Expected Units") or c.get("expected_units_counted")
        if exp is not None:
            total += int(exp)
    return (len(styles), total)


def sync_manifest_from_intrasheet(po: str, *, vendor: str | None = None) -> dict[str, Any]:
    """
    Update manifest.json from intrasheet.xlsx so they match.
    Preserves existing vendor if not overridden. Returns the written manifest.
    """
    folder = _po_folder(po)
    xlsx_path = folder / "intrasheet.xlsx"
    manifest_path = folder / "manifest.json"
    if not xlsx_path.exists():
        return {}

    cartons = get_cartons_for_po(po)
    style_count, units_total = _derive_from_intrasheet(cartons)

    existing = get_po_manifest(po)
    manifest = {
        "vendor": vendor or existing.get("vendor") or "Unknown vendor",
        "expected_style_count": style_count,
        "expected_units_total": units_total,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2))
    return manifest


def list_available_pos() -> list[str]:
    """
    Return list of PO numbers that have package folders.
    Used to identify/resolve spoken PO against known packages.
    """
    if not DATA_DIR.exists():
        return []
    pos = []
    for p in DATA_DIR.iterdir():
        if p.is_dir() and not p.name.startswith("_") and (p / "intrasheet.xlsx").exists():
            pos.append(p.name)
    return sorted(pos)


def resolve_po(spoken: str) -> str | None:
    """
    Resolve spoken PO input to an existing PO package.
    Returns the matched PO number or None if no match.
    """
    normalized = _normalize_po(spoken)
    if not normalized:
        return None
    available = list_available_pos()
    if normalized in available:
        return normalized
    # Try prefix match: "12345" might mean "123456"
    for po in available:
        if po.startswith(normalized) or normalized.startswith(po):
            return po
    return None


def get_cartons_for_po(po: str) -> list[dict[str, Any]]:
    """
    Load carton list for a PO from Excel intrasheet.
    Returns list of carton records with inspection fields.
    """
    folder = _po_folder(po)
    if not folder.exists():
        return []

    xlsx_path = folder / "intrasheet.xlsx"
    if not xlsx_path.exists():
        return []

    try:
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb.active
        if not ws:
            return []

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return []

        headers = [str(h).strip().lower().replace(" ", "_") if h else "" for h in rows[0]]
        cartons = []
        for i, row in enumerate(rows[1:], start=2):
            record = {"id": f"carton_{i}", "row": i}
            for j, val in enumerate(row):
                if j < len(headers) and headers[j]:
                    key = headers[j]
                    record[key] = val
            # Normalize keys for frontend
            if "carton_id" in record and "barcode" not in record:
                record["barcode"] = record.get("carton_id")
            cartons.append(record)
        wb.close()
        return cartons
    except Exception:
        return []


def ensure_po_folder(po: str) -> Path:
    """Create PO folder if it doesn't exist."""
    folder = _po_folder(po)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def create_intrasheet_template(po: str) -> Path:
    """
    Create a blank intrasheet.xlsx with standard inspection columns.
    """
    folder = ensure_po_folder(po)
    xlsx_path = folder / "intrasheet.xlsx"
    if xlsx_path.exists():
        return xlsx_path

    try:
        import openpyxl
        from openpyxl.styles import Font

        wb = openpyxl.Workbook()
        ws = wb.active
        if ws:
            ws.title = "Cartons"
            headers = [
                "Carton ID",
                "Barcode",
                "Style ID",
                "Expected Units",
                "Color",
                "Size",
                "Units Counted",
                "Disposition",
                "Status",
                "Notes",
            ]
            for col, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True)
            wb.save(xlsx_path)
        return xlsx_path
    except Exception:
        raise
